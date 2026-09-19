import os
import re
import openpyxl
import pandas as pd
from dotenv import load_dotenv

XLSX_PATH = r"C:\Users\kornelius\Downloads\Teacher on Duty_2026-2027.xlsx"
OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.abspath(os.path.join(OUTPUT_DIR, ".."))
CSV_PATH = os.path.join(PROJECT_ROOT, "JADWAL_DUTY_MASTER.csv")

# Standardize teacher names if short names are used
NAME_STANDARDIZATION = {
    "Kornel": "Mr. Kornelius",
    "Kornelius": "Mr. Kornelius",
    "Jenny": "Ms. Jenny P",
    "Jenny Prajitno": "Ms. Jenny P",
    "Rahayu": "Ms. Rahayu",
    "Rahayu Tjandarso": "Ms. Rahayu",
    "Minh": "Mr. Minh",
    "Minh Nguyen": "Mr. Minh",
    "Destha": "Ms. Destha",
    "Adhik": "Mr. Adhik",
    "Adhik M.Pinastiko": "Mr. Adhik",
    "Angga": "Mr. Angga",
    "Dhevangga": "Mr. Angga",
    "Kristiani": "Ms. Kristiani",
    "Gloria": "Ms. Gloria",
    "Gloria M.Rewah": "Ms. Gloria",
    "Gideon": "Mr. Gideon",
    "Albert": "Mr. Albert",
    "Albertus Christian Yuwono": "Mr. Albert",
    "Ridham": "Mr. Ridham",
    "Inggrid": "Ms. Inggrid",
    "Percy": "Mr. Percy",
    "Wang": "Ms. Wang",
    "Wang Yan": "Ms. Wang",
    "Endah": "Ms. Endah",
    "Endah Rosita": "Ms. Endah",
    "Sherly": "Ms. Sherly",
    "Sherly Puspitasari": "Ms. Sherly",
    "Shenny": "Ms. Shenny",
    "Angi": "Ms. Angi",
    "Pelangi Adonai": "Ms. Angi",
    "Rachelle": "Ms. Rachelle",
    "Meitha": "Ms. Meitha",
    "Hendro": "Mr. Hendro",
    "Hendro Suseno": "Mr. Hendro",
    "Dion": "Mr. Dion",
    "Dionisius Andrew Wibisono": "Mr. Dion",
    "Nano": "Mr. Nano",
    "Hendy": "Mr. Hendy",
    "Citra": "Ms. Citra",
    "Vita": "Ms. Vita",
    "Agnes": "Ms. Agnes",
    "Phoebe": "Ms. Phoebe",
    "Hindra": "Mr. Hindra",
    "Hakim": "Mr. Hakim",
    "Ardhi": "Mr. Ardhi",
    "Benu": "Mr. Benu",
    "Sus": "Ms. Sus"
}

def standardize_name(raw_name):
    clean = str(raw_name).strip()
    return NAME_STANDARDIZATION.get(clean, clean)

def parse_and_sync_duty():
    print(f"Reading duty schedule from: {XLSX_PATH}...")
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    
    duty_records = []
    days = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday']

    # 1. Morning Duty
    s_morn = wb['7.15-7.45']
    for r in range(4, 9):
        loc = str(s_morn.cell(r, 1).value or '').strip()
        task = str(s_morn.cell(r, 7).value or '').strip()
        for col_idx, d in enumerate(days, start=2):
            raw_t = str(s_morn.cell(r, col_idx).value or '').strip()
            if raw_t:
                duty_records.append({
                    'category': 'Morning Duty',
                    'grade_scope': 'All Grades (Schoolwide)',
                    'location': loc,
                    'day_of_week': d,
                    'time_slot': '07.15-07.45',
                    'teacher_name': standardize_name(raw_t),
                    'task': task
                })

    # 2. Gr 1-2
    s_gr12 = wb['Gr 1-2']
    sections_12 = [
        ('Break 1', [4, 5], {'Monday': '09.10-09.35', 'Tuesday': '09.10-09.35', 'Wednesday': '09.10-09.35', 'Thursday': '09.30-09.45', 'Friday': '09.10-09.35'}),
        ('Break 2', [7, 8], {'Monday': '11.20-11.50', 'Tuesday': '11.20-11.50', 'Wednesday': '11.20-11.50', 'Thursday': '11.30-11.50', 'Friday': '11.20-11.50'}),
        ('Go Home', [10, 11], {'Monday': '13.00-13.35', 'Tuesday': '13.00-13.35', 'Wednesday': '13.00-13.35', 'Thursday': '13.25-13.50', 'Friday': '13.00-13.35'})
    ]
    for cat, rows, time_map in sections_12:
        for r in rows:
            loc = str(s_gr12.cell(r, 1).value or '').strip()
            task = str(s_gr12.cell(r, 7).value or '').strip()
            for col_idx, d in enumerate(days, start=2):
                raw_t = str(s_gr12.cell(r, col_idx).value or '').strip()
                if raw_t:
                    duty_records.append({
                        'category': f'{cat} (Grade 1-2)',
                        'grade_scope': 'Grade 1-2',
                        'location': loc,
                        'day_of_week': d,
                        'time_slot': time_map[d],
                        'teacher_name': standardize_name(raw_t),
                        'task': task
                    })

    # 3. Gr 3-4
    s_gr34 = wb['Gr 3-4']
    sections_34 = [
        ('Break 1', [4, 5, 6], {'Monday': '09.10-09.35', 'Tuesday': '09.10-09.35', 'Wednesday': '09.30-09.45', 'Thursday': '09.10-09.35', 'Friday': '09.10-09.35'}),
        ('Break 2', [8, 9, 10], {'Monday': '11.55-12.25', 'Tuesday': '11.55-12.25', 'Wednesday': '11.30-11.50', 'Thursday': '11.55-12.25', 'Friday': '11.55-12.25'}),
        ('Go Home', [12, 13], {'Monday': '13.35-14.00', 'Tuesday': '13.35-14.00', 'Wednesday': '13.50-14.00', 'Thursday': '13.35-14.00', 'Friday': '13.35-14.00'})
    ]
    for cat, rows, time_map in sections_34:
        for r in rows:
            loc = str(s_gr34.cell(r, 1).value or '').strip()
            task = str(s_gr34.cell(r, 7).value or '').strip()
            for col_idx, d in enumerate(days, start=2):
                raw_t = str(s_gr34.cell(r, col_idx).value or '').strip()
                if raw_t:
                    duty_records.append({
                        'category': f'{cat} (Grade 3-4)',
                        'grade_scope': 'Grade 3-4',
                        'location': loc,
                        'day_of_week': d,
                        'time_slot': time_map[d],
                        'teacher_name': standardize_name(raw_t),
                        'task': task
                    })

    # 4. Gr 5-6
    s_gr56 = wb['Gr 5-6']
    sections_56 = [
        ('Break 1', [4, 5, 6, 7], {'Monday': '10.20-10.45', 'Tuesday': '10.20-10.45', 'Wednesday': '10.20-10.45', 'Thursday': '10.20-10.45', 'Friday': '10.20-10.45'}),
        ('Break 2', [9, 10, 11, 12], {'Monday': '11.55-12.25', 'Tuesday': '11.55-12.25', 'Wednesday': '11.55-12.25', 'Thursday': '12.05-12.25', 'Friday': '11.55-12.25'}),
        ('Go Home', [14, 15], {'Monday': '14.10-14.35', 'Tuesday': '14.10-14.35', 'Wednesday': '14.10-14.35', 'Thursday': '14.35-14.50', 'Friday': '14.10-14.35'})
    ]
    for cat, rows, time_map in sections_56:
        for r in rows:
            loc = str(s_gr56.cell(r, 1).value or '').strip()
            task = str(s_gr56.cell(r, 7).value or '').strip()
            for col_idx, d in enumerate(days, start=2):
                raw_t = str(s_gr56.cell(r, col_idx).value or '').strip()
                if raw_t:
                    duty_records.append({
                        'category': f'{cat} (Grade 5-6)',
                        'grade_scope': 'Grade 5-6',
                        'location': loc,
                        'day_of_week': d,
                        'time_slot': time_map[d],
                        'teacher_name': standardize_name(raw_t),
                        'task': task
                    })

    # Save to CSV
    df = pd.DataFrame(duty_records)
    df.to_csv(CSV_PATH, index=False, encoding="utf-8")
    print(f"Saved {len(df)} duty records to: {CSV_PATH}")

    # Sync to MySQL database
    load_dotenv(os.path.join(OUTPUT_DIR, ".env"))
    import database
    import models

    # Create table if not exists
    models.Base.metadata.create_all(bind=database.engine)

    db = next(database.get_db())
    admin_obj = db.query(models.Admin).first()
    admin_id = admin_obj.id_admin if admin_obj else 1

    try:
        deleted = db.query(models.TeacherDuty).delete()
        print(f"Deleted {deleted} previous duty rows from MySQL.")
        
        db_items = [
            models.TeacherDuty(
                category=r['category'],
                grade_scope=r['grade_scope'],
                location=r['location'],
                day_of_week=r['day_of_week'],
                time_slot=r['time_slot'],
                teacher_name=r['teacher_name'],
                task=r['task'],
                admin_update=admin_id
            )
            for _, r in df.iterrows()
        ]
        db.bulk_save_objects(db_items)
        db.commit()
        print(f"Successfully inserted {len(db_items)} duty records into teacher_duties table!")
    except Exception as e:
        db.rollback()
        print("Database sync error:", e)
        raise
    finally:
        db.close()

    return len(df)

if __name__ == "__main__":
    parse_and_sync_duty()
