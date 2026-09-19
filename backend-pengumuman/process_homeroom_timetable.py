import os
import re
import openpyxl
import pandas as pd

XLSX_PATH = r"C:\Users\kornelius\Downloads\CLASS_Time Table 26-27.xlsx"
OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.abspath(os.path.join(OUTPUT_DIR, ".."))

OUT_MODIFIED_CSV = os.path.join(PROJECT_ROOT, "JADWAL_KELAS_MASTER_MODIFIED.csv")
OUT_RELATIONAL_CSV = os.path.join(PROJECT_ROOT, "JADWAL_KELAS_MASTER.csv")
EXISTING_MASTER_CSV = os.path.join(PROJECT_ROOT, "JADWAL_GURU_MASTER.csv")
COMBINED_MASTER_CSV = os.path.join(PROJECT_ROOT, "JADWAL_GURU_MASTER.csv")

def is_homeroom_subject(grade_str, subject_name):
    s = subject_name.strip()
    m = re.search(r'\d', str(grade_str))
    g = int(m.group()) if m else 1
    
    # Common homeroom subjects across all grades
    if s in ['PC Session', 'Assembly', 'UoI', 'BI', 'P.Pancasila', 'Preparation', 'Break', 'Excur']:
        return True
    # Math is taught by Homeroom only in Grade 1, 2, 3, 4
    if s in ['Math'] and g in [1, 2, 3, 4]:
        return True
    return False

def process_homeroom_timetable():
    print(f"Loading Excel file from: {XLSX_PATH}...")
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    
    # 1. Load specialist lookup from existing master if available
    spec_lookup = {}
    if os.path.exists(EXISTING_MASTER_CSV):
        df_old = pd.read_csv(EXISTING_MASTER_CSV)
        df_spec_clean = df_old[~df_old['subject_grade'].str.startswith("Homeroom", na=False)]
        for _, r in df_spec_clean.iterrows():
            c = str(r['class_name']).strip()
            d = str(r['day_of_week']).strip()
            s = str(r['time_slot']).strip()
            spec_lookup[(d, c, s)] = (str(r['teacher_name']).strip(), str(r['subject_grade']).strip())
    
    all_homeroom_teacher_rows = []
    all_class_timetable_rows = []
    matrix_rows = []
    
    grade_sheets = ['Grade 1', 'Grade 2', 'Grade 3', 'Grade 4', 'Grade 5', 'Grade 6']
    
    for g in grade_sheets:
        sheet = wb[g]
        print(f"Processing sheet: {g}...")
        
        # Unmerge and propagate values
        unmerged = {}
        for r in range(1, sheet.max_row + 1):
            for c in range(1, sheet.max_column + 1):
                unmerged[(r, c)] = sheet.cell(r, c).value

        for rng in sheet.merged_cells.ranges:
            min_col, min_row, max_col, max_row = rng.min_col, rng.min_row, rng.max_col, rng.max_row
            top_val = unmerged.get((min_row, min_col))
            for r in range(min_row, max_row + 1):
                for c in range(min_col, max_col + 1):
                    unmerged[(r, c)] = top_val

        # Identify class headers (e.g. '1A (Ms. Debby and Ms. Destha )')
        class_headers = []
        for r in range(1, sheet.max_row + 1):
            val = str(unmerged.get((r, 1)) or '').strip()
            m = re.match(r'^([1-6][A-Ea-e])\s*\(([^)]+)\)', val)
            if m:
                class_headers.append((r, m.group(1).upper(), m.group(2).strip(), val))

        is_gr34 = g in ['Grade 3', 'Grade 4']

        for idx, (r_start, cls_name, teachers_str, raw_header) in enumerate(class_headers):
            r_end = class_headers[idx + 1][0] - 1 if idx + 1 < len(class_headers) else sheet.max_row
            
            t_parts = re.split(r'\s+and\s+|\s*,\s*', teachers_str)
            main_teacher = t_parts[0].strip() if len(t_parts) > 0 else teachers_str

            # Matrix output header
            if is_gr34:
                matrix_header = ["Time 1", "Monday", "Tuesday", "Time 2", "Wednesday", "Time 1", "Thursday", "Friday"]
            else:
                matrix_header = ["Time 1", "Monday", "Tuesday", "Wednesday", "Time 2", "Thursday", "Time 1", "Friday"]

            matrix_rows.append({col: "" for col in matrix_header})
            matrix_rows.append({matrix_header[0]: f"{cls_name} ({teachers_str})"})
            matrix_rows.append({col: col for col in matrix_header})

            # Read rows inside class block
            for r in range(r_start + 2, r_end + 1):
                if is_gr34:
                    t1 = str(unmerged.get((r, 1)) or '').strip()
                    mon = str(unmerged.get((r, 2)) or '').strip()
                    tue = str(unmerged.get((r, 3)) or '').strip()
                    t2 = str(unmerged.get((r, 4)) or '').strip()
                    wed = str(unmerged.get((r, 5)) or '').strip()
                    thu = str(unmerged.get((r, 6)) or '').strip()
                    fri = str(unmerged.get((r, 7)) or '').strip()

                    if not any([t1, mon, tue, t2, wed, thu, fri]):
                        continue

                    matrix_rows.append({
                        "Time 1": t1, "Monday": mon, "Tuesday": tue,
                        "Time 2": t2, "Wednesday": wed, "Time 1_dup": t1,
                        "Thursday": thu, "Friday": fri
                    })

                    day_slots = [
                        ("Monday", t1, mon), ("Tuesday", t1, tue),
                        ("Wednesday", t2, wed), ("Thursday", t2, thu), ("Friday", t2, fri)
                    ]
                else:
                    t1 = str(unmerged.get((r, 1)) or '').strip()
                    mon = str(unmerged.get((r, 2)) or '').strip()
                    tue = str(unmerged.get((r, 3)) or '').strip()
                    wed = str(unmerged.get((r, 4)) or '').strip()
                    t2 = str(unmerged.get((r, 5)) or '').strip()
                    thu = str(unmerged.get((r, 6)) or '').strip()
                    fri = str(unmerged.get((r, 7)) or '').strip()

                    if not any([t1, mon, tue, wed, t2, thu, fri]):
                        continue

                    matrix_rows.append({
                        "Time 1": t1, "Monday": mon, "Tuesday": tue, "Wednesday": wed,
                        "Time 2": t2, "Thursday": thu, "Time 1_dup": t1, "Friday": fri
                    })

                    day_slots = [
                        ("Monday", t1, mon), ("Tuesday", t1, tue), ("Wednesday", t1, wed),
                        ("Thursday", t2, thu), ("Friday", t2, fri)
                    ]

                for d_name, t_slot, subj in day_slots:
                    if not t_slot or '-' not in t_slot or not subj:
                        continue

                    # 1. Full Class Timetable entry
                    # Determine who actually teaches this slot
                    if is_homeroom_subject(g, subj):
                        actual_teacher = main_teacher if subj != "Break" else "Break"
                        teacher_role = "Wali Kelas"
                    else:
                        # Specialist lesson
                        spec_info = spec_lookup.get((d_name, cls_name, t_slot))
                        if spec_info:
                            actual_teacher = spec_info[0]
                            teacher_role = "Guru Spesialis"
                        else:
                            actual_teacher = f"Guru Spesialis ({subj})"
                            teacher_role = "Guru Spesialis"

                    all_class_timetable_rows.append({
                        "class_name": cls_name,
                        "day_of_week": d_name,
                        "time_slot": t_slot,
                        "subject": subj,
                        "teacher_name": actual_teacher,
                        "teacher_role": teacher_role,
                        "grade": g
                    })

                    # 2. Homeroom Teacher's own teaching schedule entry
                    # Homeroom teacher DOES NOT teach specialist subjects!
                    subj_grade = f"Homeroom {g} ({cls_name})"
                    if is_homeroom_subject(g, subj):
                        hr_note = subj
                    else:
                        hr_note = f"Free (Sesi {subj})"

                    all_homeroom_teacher_rows.append({
                        "subject_grade": subj_grade,
                        "teacher_name": main_teacher,
                        "day_of_week": d_name,
                        "time_slot": t_slot,
                        "class_name": cls_name,
                        "note": hr_note
                    })

    # Save JADWAL_KELAS_MASTER.csv (Full Class Timetable: homeroom + specialist)
    df_class_full = pd.DataFrame(all_class_timetable_rows)
    df_class_full.to_csv(OUT_RELATIONAL_CSV, index=False, encoding="utf-8")
    print(f"Saved {len(df_class_full)} full class timetable records to: {OUT_RELATIONAL_CSV}")

    # Save JADWAL_KELAS_MASTER_MODIFIED.csv
    df_matrix = pd.DataFrame(matrix_rows)
    df_matrix.to_csv(OUT_MODIFIED_CSV, index=False, encoding="utf-8")
    print(f"Saved {len(df_matrix)} matrix rows to: {OUT_MODIFIED_CSV}")

    # Save Combined JADWAL_GURU_MASTER.csv (Specialist Teachers + Homeroom Teachers with non-specialist teaching slots)
    df_homeroom = pd.DataFrame(all_homeroom_teacher_rows)
    if os.path.exists(EXISTING_MASTER_CSV):
        df_specialist = pd.read_csv(EXISTING_MASTER_CSV)
        df_specialist_clean = df_specialist[~df_specialist['subject_grade'].str.startswith("Homeroom", na=False)]
        
        df_combined = pd.concat([df_specialist_clean, df_homeroom], ignore_index=True)
        df_combined.to_csv(COMBINED_MASTER_CSV, index=False, encoding="utf-8")
        print(f"Combined master updated: {len(df_specialist_clean)} specialist + {len(df_homeroom)} homeroom = {len(df_combined)} total records in {COMBINED_MASTER_CSV}")
        return len(df_combined)
    return len(df_homeroom)

if __name__ == "__main__":
    process_homeroom_timetable()
